"""Step 22 Part 16: export reproducibility. Generates every export
format twice from ONE frozen, on-disk `SqliteValidationStore`, and
proves every quantitative value is identical between the two runs
(only generation timestamps/file bytes may differ), then cross-checks
the exported figures directly against the source records -- no
reporting layer independently recalculates an authoritative value.
"""
from __future__ import annotations

import csv
import json
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

import pytest
from openpyxl import load_workbook

from src.reporting.export import export_validation_cohort
from src.validation.session import DailySnapshot, RuleViolationRecord, SqliteValidationStore

from .test_persistence_restart_recovery import _build_real_opportunity, _cohort_record, _trade_record

NOW = datetime(2026, 9, 22, 14, 0, tzinfo=timezone.utc)


async def _populated_store(tmp_path: Path) -> SqliteValidationStore:
    db_path = tmp_path / "export_test.db"
    store = SqliteValidationStore(db_path)
    cohort = _cohort_record("cohort-export-1")
    store.record_cohort(cohort)
    store.record_snapshot(
        DailySnapshot(snapshot_date=date(2026, 9, 22), nav=100_000.0, cash=95_000.0, capital_deployed_pct=0.05,
                       open_position_count=0, drawdown_pct=0.0, per_strategy_nav={}, recorded_at=NOW),
        cohort_id="cohort-export-1",
    )
    store.record_snapshot(
        DailySnapshot(snapshot_date=date(2026, 9, 23), nav=99_200.0, cash=94_500.0, capital_deployed_pct=0.06,
                       open_position_count=1, drawdown_pct=0.008, per_strategy_nav={"put_credit_spread": -800.0}, recorded_at=NOW + timedelta(days=1)),
        cohort_id="cohort-export-1",
    )
    opportunity = await _build_real_opportunity("cohort-export-1")
    store.record_opportunity(opportunity)
    store.record_trade(_trade_record())
    store.record_violation(RuleViolationRecord(
        violation_id="viol-export-1", occurred_at=NOW, rule_name="test_rule", description="exercised for export reproducibility",
    ))
    return store


class TestExportReproducibility:
    @pytest.mark.asyncio
    async def test_two_export_runs_from_the_same_frozen_store_produce_identical_quantitative_content(self, tmp_path):
        store = await _populated_store(tmp_path)
        out_1 = tmp_path / "export_run_1"
        out_2 = tmp_path / "export_run_2"

        manifest_1 = export_validation_cohort(store, cohort_id="cohort-export-1", output_dir=out_1, generated_at=NOW)
        manifest_2 = export_validation_cohort(store, cohort_id="cohort-export-1", output_dir=out_2, generated_at=NOW + timedelta(hours=1))

        # --- JSON: identical except the top-level generated_at ---
        json_1 = json.loads(manifest_1.json_file.read_text())
        json_2 = json.loads(manifest_2.json_file.read_text())
        json_1.pop("generated_at")
        json_2.pop("generated_at")
        assert json_1 == json_2

        # --- CSV: byte-identical (no timestamp embedded in any of these) ---
        for name, path_1 in manifest_1.csv_files.items():
            path_2 = manifest_2.csv_files[name]
            assert path_1.read_text() == path_2.read_text(), f"{name}.csv differs between two runs from the same store"

        # --- XLSX: identical cell values on every sheet except the
        # Summary sheet's own "Generated At" row ---
        wb_1 = load_workbook(manifest_1.xlsx_file)
        wb_2 = load_workbook(manifest_2.xlsx_file)
        assert wb_1.sheetnames == wb_2.sheetnames
        for sheet_name in wb_1.sheetnames:
            ws_1, ws_2 = wb_1[sheet_name], wb_2[sheet_name]
            rows_1 = [[c.value for c in row] for row in ws_1.iter_rows()]
            rows_2 = [[c.value for c in row] for row in ws_2.iter_rows()]
            if sheet_name == "Summary":
                rows_1 = [r for r in rows_1 if not (r and r[0] == "Generated At")]
                rows_2 = [r for r in rows_2 if not (r and r[0] == "Generated At")]
            assert rows_1 == rows_2, f"XLSX sheet {sheet_name!r} differs between two runs"

        # --- PDF: both were generated successfully and are non-trivial files ---
        assert manifest_1.pdf_file.exists() and manifest_1.pdf_file.stat().st_size > 500
        assert manifest_2.pdf_file.exists() and manifest_2.pdf_file.stat().st_size > 500

    @pytest.mark.asyncio
    async def test_exported_figures_cross_check_against_the_source_database_values(self, tmp_path):
        """Specifically verifies the figures Part 16 names by name:
        NAV, trade count, win rate, and P&L are traced from the export
        straight back to the exact source records -- never a
        reporting-layer recalculation that could drift."""
        store = await _populated_store(tmp_path)
        out = tmp_path / "cross_check_export"
        manifest = export_validation_cohort(store, cohort_id="cohort-export-1", output_dir=out, generated_at=NOW)

        source_trades = store.trades()
        source_snapshots = store.snapshots(cohort_id="cohort-export-1")
        assert len(source_trades) == 1
        source_trade = source_trades[0]

        # NAV: the CSV's own daily_portfolio.csv values match the exact
        # DailySnapshot.nav figures recorded in the store.
        with manifest.csv_files["daily_portfolio"].open() as f:
            csv_navs = {row["snapshot_date"]: float(row["nav"]) for row in csv.DictReader(f)}
        for s in source_snapshots:
            assert csv_navs[s.snapshot_date.isoformat()] == s.nav

        # Trade count / win rate / P&L: the strategy_performance.csv
        # row for this trade's strategy matches a hand-computed
        # aggregate over the exact same source TradeRecord.
        with manifest.csv_files["strategy_performance"].open() as f:
            strat_rows = {row["strategy"]: row for row in csv.DictReader(f)}
        row = strat_rows[source_trade.strategy.value]
        assert int(row["trade_count"]) == 1
        assert float(row["win_rate"]) == (1.0 if source_trade.realistic_pnl > 0 else 0.0)
        assert float(row["total_realistic_pnl"]) == pytest.approx(source_trade.realistic_pnl)
        assert float(row["expectancy"]) == pytest.approx(source_trade.realistic_pnl)  # single trade -- expectancy == its own P&L

        # The JSON export's trade record matches the source TradeRecord exactly.
        json_payload = json.loads(manifest.json_file.read_text())
        assert len(json_payload["trades"]) == 1
        assert json_payload["trades"][0]["position_id"] == source_trade.position_id
        assert json_payload["trades"][0]["realistic_pnl"] == pytest.approx(source_trade.realistic_pnl)

        # The bundle itself (what every format was built from) is the
        # single source every cross-check above ultimately traces to.
        bundle = manifest.bundle
        assert bundle.trades == tuple(source_trades)
        assert bundle.risk_metrics.starting_nav == source_snapshots[0].nav if source_snapshots else True

    @pytest.mark.asyncio
    async def test_export_schema_version_is_present_and_stable_across_runs(self, tmp_path):
        store = await _populated_store(tmp_path)
        out_1 = tmp_path / "version_run_1"
        out_2 = tmp_path / "version_run_2"
        m1 = export_validation_cohort(store, cohort_id="cohort-export-1", output_dir=out_1, generated_at=NOW)
        m2 = export_validation_cohort(store, cohort_id="cohort-export-1", output_dir=out_2, generated_at=NOW)
        assert m1.bundle.export_schema_version == m2.bundle.export_schema_version
        assert m1.bundle.export_schema_version  # non-empty


class TestExportsHandleAnEmptyCohortHonestly:
    """No fabricated numbers for a cohort with nothing recorded yet --
    every section reports zero/empty/N/A explicitly, never a made-up
    placeholder value."""

    def test_empty_store_exports_cleanly_with_zeroed_not_fabricated_metrics(self, tmp_path):
        db_path = tmp_path / "empty.db"
        store = SqliteValidationStore(db_path)
        cohort = _cohort_record("empty-cohort")
        store.record_cohort(cohort)

        out = tmp_path / "empty_export"
        manifest = export_validation_cohort(store, cohort_id="empty-cohort", output_dir=out, generated_at=NOW)

        assert manifest.bundle.trades == ()
        assert manifest.bundle.risk_metrics.starting_nav is None
        assert manifest.bundle.risk_metrics.max_drawdown_pct is None
        assert manifest.pdf_file.exists()
        assert manifest.xlsx_file.exists()
        assert manifest.json_file.exists()

        payload = json.loads(manifest.json_file.read_text())
        assert payload["trades"] == []
        assert "N/A" in payload["benchmark_comparison"]
