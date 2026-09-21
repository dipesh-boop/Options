"""Part 12: the XLSX validation workbook -- one worksheet per section,
quantitative values stored as real numeric cells (never
formatted-as-text), generation timestamp and cohort ID on the Summary
sheet.
"""
from __future__ import annotations

import math
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from src.reporting.data import ValidationReportBundle


def _num(value):
    """openpyxl cannot store `inf`/`nan` as a numeric cell (raises) --
    an unbounded max_profit is written as the literal string "inf",
    exactly as honestly labeled as the platform's own JSON export
    (never silently coerced to some large finite number)."""
    if isinstance(value, float) and (math.isinf(value) or math.isnan(value)):
        return "inf" if value > 0 else ("-inf" if value < 0 else "nan")
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value


def _write_table(ws: Worksheet, headers: list[str], rows: list[dict]) -> None:
    ws.append(headers)
    for row in rows:
        ws.append([_num(row.get(h)) for h in headers])


def export_xlsx(bundle: ValidationReportBundle, output_path: Path | str) -> Path:
    wb = Workbook()

    summary = wb.active
    summary.title = "Summary"
    summary.append(["Field", "Value"])
    summary.append(["Generated At", bundle.generated_at.isoformat()])
    summary.append(["Export Schema Version", bundle.export_schema_version])
    summary.append(["Cohort ID", bundle.cohort.cohort_id if bundle.cohort else "(none)"])
    summary.append(["Cohort Name", bundle.cohort.cohort_name if bundle.cohort else ""])
    summary.append(["Cohort Status", bundle.cohort.status if bundle.cohort else ""])
    summary.append(["Manifest ID", bundle.cohort.manifest.manifest_id if bundle.cohort else ""])
    summary.append(["Trade Count", len(bundle.trades)])
    summary.append(["Snapshot Count", len(bundle.snapshots)])
    summary.append(["Opportunity Count", len(bundle.opportunities)])
    summary.append(["Rule Violation Count", len(bundle.violations)])
    summary.append(["Unresolved Reconciliation Failures", bundle.risk_metrics.unresolved_reconciliation_failure_count])

    daily_nav = wb.create_sheet("Daily NAV")
    _write_table(daily_nav, ["snapshot_date", "nav", "cash", "capital_deployed_pct", "open_position_count", "drawdown_pct"], [
        {"snapshot_date": s.snapshot_date, "nav": s.nav, "cash": s.cash, "capital_deployed_pct": s.capital_deployed_pct,
         "open_position_count": s.open_position_count, "drawdown_pct": s.drawdown_pct}
        for s in sorted(bundle.snapshots, key=lambda s: s.snapshot_date)
    ])

    trades_ws = wb.create_sheet("Trades")
    _write_table(trades_ws, [
        "position_id", "ticker", "strategy", "contracts", "opened_at", "closed_at", "close_reason",
        "capital_at_risk", "realistic_pnl", "theoretical_pnl", "commission_paid",
    ], [
        {"position_id": t.position_id, "ticker": t.ticker, "strategy": t.strategy.value, "contracts": t.contracts,
         "opened_at": t.opened_at, "closed_at": t.closed_at, "close_reason": t.close_reason,
         "capital_at_risk": t.capital_at_risk, "realistic_pnl": t.realistic_pnl,
         "theoretical_pnl": t.theoretical_pnl, "commission_paid": t.commission_paid}
        for t in sorted(bundle.trades, key=lambda t: t.opened_at)
    ])

    strategies_ws = wb.create_sheet("Strategies")
    _write_table(strategies_ws, [
        "strategy", "trade_count", "win_count", "win_rate", "total_realistic_pnl", "expectancy", "largest_win", "largest_loss",
    ], [vars(r) for r in bundle.strategy_performance])

    risk_ws = wb.create_sheet("Risk")
    rm = bundle.risk_metrics
    risk_ws.append(["Metric", "Value"])
    for key, value in vars(rm).items():
        risk_ws.append([key, _num(value)])

    execution_ws = wb.create_sheet("Execution")
    _write_table(execution_ws, [
        "strategy", "trade_count", "average_entry_spread_pct", "average_realistic_vs_theoretical_entry_gap", "total_commission_paid",
    ], [vars(r) for r in bundle.execution_quality])

    benchmarks_ws = wb.create_sheet("Benchmarks")
    benchmarks_ws.append(["Note"])
    benchmarks_ws.append([bundle.benchmark_comparison_note])

    decision_ws = wb.create_sheet("Decision Quality")
    _write_table(decision_ws, ["opportunity_id", "ticker", "pipeline_status", "risk_decision", "cash_no_trade", "alternative_count"],
                 [vars(r) for r in bundle.decision_quality])

    counterfactuals_ws = wb.create_sheet("Counterfactuals")
    _write_table(counterfactuals_ws, [
        "opportunity_id", "strategy_kind", "was_selected", "selection_or_rejection_reason",
        "expected_value", "maximum_loss", "probability_of_profit",
    ], [vars(r) for r in bundle.selection_records])

    compliance_ws = wb.create_sheet("Rule Compliance")
    _write_table(compliance_ws, ["violation_id", "occurred_at", "rule_name", "description", "related_id"], [
        {"violation_id": v.violation_id, "occurred_at": v.occurred_at, "rule_name": v.rule_name,
         "description": v.description, "related_id": v.related_id}
        for v in bundle.violations
    ])

    configuration_ws = wb.create_sheet("Configuration")
    if bundle.cohort is not None:
        configuration_ws.append(["Config File", "SHA-256 Hash"])
        for path, file_hash in bundle.cohort.manifest.config_file_hashes.items():
            configuration_ws.append([path, file_hash])
        configuration_ws.append([])
        configuration_ws.append(["Strategy/Component Version Key", "Value"])
        for key, value in bundle.cohort.manifest.strategy_versions.items():
            configuration_ws.append([key, value])

    manifest_ws = wb.create_sheet("Manifest")
    if bundle.cohort is not None:
        m = bundle.cohort.manifest
        manifest_ws.append(["Field", "Value"])
        manifest_ws.append(["manifest_id", m.manifest_id])
        manifest_ws.append(["frozen_at", m.frozen_at.isoformat()])
        manifest_ws.append(["period_start", m.period.start_date.isoformat()])
        manifest_ws.append(["period_end", m.period.end_date.isoformat()])
        manifest_ws.append(["duration_days", m.period.duration_days])
        manifest_ws.append(["starting_nav", m.starting_nav])
        manifest_ws.append(["cohort_label", m.cohort_label])
        manifest_ws.append(["notes", m.notes])

    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
